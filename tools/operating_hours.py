"""日中（9〜18時）の室外機の稼働時間を日ごとに数え、Excel とスライド用のデータにまとめる。

目的
    研究に使える時間（室外機が運転している時間）が、7/29以降どう変わったかを日ごとに見える化する。

入力
    実測データの Excel（シート「整理済みデータ」、1分値）
    例：python3 tools/operating_hours.py 261013_実測データ_完全版.xlsx 出力.xlsx 出力.json

処理（運転の条件は tools/operation_judge.py と同じ）
    1. 1分ごと：排気温度 − 吸気温度（3方位の平均）の10分移動中央値が1.0℃以上なら「運転中」
    2. 1時間ごと：判定できた分が48分以上あり、そのうち80％以上が運転中なら「運転」
    3. 日ごと：9〜17時台（9:00〜18:00）の9時間のうち「運転」の時間を数える。
       判定できた時間が7時間未満の日は空欄（欠測で少なく数えないため）
    4. 列の値は3台の平均（その日に値がある室外機だけで平均）
    5. 研究に使える時間：平日に、組（芋緑化ありの室外機と同じ区画の列3）の2台がともに「運転」の時間

出力
    Excel（日ごとの稼働時間、研究に使える時間、期間別のまとめ、グラフ、条件）と、スライド用の JSON
"""
import json
import sys

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, __file__.rsplit("/", 1)[0])
import operation_judge as oj  # noqa: E402

DAY_START, DAY_END = 9, 17           # 9時台〜17時台
MIN_JUDGED = 7                       # その日に判定できた時間がこれ未満なら空欄
COLS = {"列3": ["3-1", "3-2", "3-3"], "列8": ["8-1", "8-2", "8-3"], "列11": ["11-1", "11-2", "11-3"]}
PERIODS = [("8月", "2026-07-30", "2026-08-31"), ("9月前半", "2026-09-01", "2026-09-15"),
           ("9月後半", "2026-09-16", "2026-09-30"), ("10月", "2026-10-01", "2026-10-13")]
COLOR = {"列3": "5E6B7A", "列8": "01B8AA", "列11": "118DFF"}
UNIT_COLOR = {"3-1": "22354F", "3-2": "5E6B7A", "3-3": "A7B1BC",
              "8-1": "00786E", "8-2": "01B8AA", "8-3": "7FD8CE",
              "11-1": "0B5CAD", "11-2": "118DFF", "11-3": "8CC4FF"}


def compute(src):
    df = oj.load(src)
    c = oj.hourly_counts(oj.diff(df))
    cls = {}
    for u in oj.UNITS:
        v, n = c[(u, "有効")], c[(u, "運転")]
        cls[u] = pd.Series(np.where(v < 48, "判定不能", np.where(n >= 0.8 * v, "運転",
                                                              np.where(n <= 0.2 * v, "停止", "断続"))), index=c.index)
    cls = pd.DataFrame(cls)
    day = cls[(cls.index.hour >= DAY_START) & (cls.index.hour <= DAY_END)]
    g = day.index.normalize()
    days = pd.date_range(str(oj.FIRST_DAY), str(oj.LAST_DAY))
    run = (day == "運転").groupby(g).sum().reindex(days)
    judged = (day != "判定不能").groupby(g).sum().reindex(days)
    unit = run.where(judged >= MIN_JUDGED)
    col = pd.DataFrame({k: unit[v].mean(axis=1) for k, v in COLS.items()})
    kind = pd.Series([oj.kind(d.date()) for d in days], index=days)
    pair = {}
    for a, b in oj.PAIRS:
        both = ((day[a] == "運転") & (day[b] == "運転")).groupby(g).sum().reindex(days)
        ok = (judged[a] >= MIN_JUDGED) & (judged[b] >= MIN_JUDGED) & (kind == "平日")
        pair[f"{a}と{b}"] = both.where(ok)
    pair = pd.DataFrame(pair)
    paircol = pd.DataFrame({"列8": pair[[f"8-{k}と3-{k}" for k in "123"]].mean(axis=1),
                            "列11": pair[[f"11-{k}と3-{k}" for k in "123"]].mean(axis=1)})
    return days, kind, unit, col, pair, paircol


# ---------------------------------------------------------------- Excel
FONT = "游ゴシック"
F = Font(name=FONT, size=9, color="172033")
FB = Font(name=FONT, size=10, color="172033", bold=True)
FH = Font(name=FONT, size=9, color="FFFFFF", bold=True)
FN = Font(name=FONT, size=9, color="5E6B7A")
HEAD = PatternFill("solid", fgColor="22354F")
SUB = PatternFill("solid", fgColor="EEF2F6")
BOX = Border(*(Side(style="thin", color="C9D1DB"),) * 4)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
WEEK = "月火水木金土日"


def put(ws, r, c, v, font=F, fmt=None, fill=None, align=CEN):
    x = ws.cell(row=r, column=c, value=v)
    x.font, x.alignment, x.border = font, align, BOX
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    return x


def val(v):
    return None if v is None or (isinstance(v, float) and np.isnan(v)) else round(float(v), 2)


def table_sheet(ws, title, days, kind, frames, note):
    ws["A1"] = title
    ws["A1"].font = FB
    ws["A2"] = note
    ws["A2"].font = FN
    hdr = ["日付", "表示用", "曜日", "区分"] + [n for fr in frames for n in fr.columns] + ["土日祝（グラフの帯）"]
    for j, h in enumerate(hdr):
        x = put(ws, 3, 1 + j, h, FH, fill=HEAD)
        ws.column_dimensions[x.column_letter].width = 9 if j != 1 else 7
    ws.row_dimensions[3].height = 30
    for i, d in enumerate(days):
        r = 4 + i
        sub = SUB if kind[d] != "平日" else None
        put(ws, r, 1, d.date(), fmt="m/d", fill=sub)
        put(ws, r, 2, f"{d.month}/{d.day}", fill=sub)
        put(ws, r, 3, WEEK[d.weekday()], fill=sub)
        put(ws, r, 4, oj.HOLIDAYS.get(d.date(), kind[d]), fill=sub)
        j = 5
        for fr in frames:
            for n in fr.columns:
                put(ws, r, j, val(fr.at[d, n]), fmt="0.0", fill=sub)
                j += 1
        put(ws, r, j, 9 if kind[d] != "平日" else None, fill=sub)
    ws.freeze_panes = "E4"
    return 4, 3 + len(days)


def combo(ws, r0, r1, series_cols, colors, title, cat_col=2, band_col=None, width=30, height=8):
    bar = BarChart()
    bar.type, bar.grouping, bar.gapWidth = "col", "clustered", 0
    s = Series(Reference(ws, min_col=band_col, min_row=r0, max_row=r1), title="土日祝")
    s.graphicalProperties.solidFill = "E9ECF1"
    s.graphicalProperties.line.noFill = True
    bar.series.append(s)
    line = LineChart()
    for (col, name), color in zip(series_cols, colors):
        s = Series(Reference(ws, min_col=col, min_row=r0, max_row=r1), title=name)
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 19050
        s.marker.symbol = "none"
        s.smooth = False
        line.series.append(s)
    bar.set_categories(Reference(ws, min_col=cat_col, min_row=r0, max_row=r1))
    line.set_categories(Reference(ws, min_col=cat_col, min_row=r0, max_row=r1))
    bar.title = title
    bar.y_axis.title = "時間（9時間中）"
    bar.y_axis.scaling.min, bar.y_axis.scaling.max, bar.y_axis.majorUnit = 0, 9, 3
    bar.x_axis.tickLblSkip = 7
    bar.display_blanks = "gap"
    bar += line
    bar.legend.position = "b"
    bar.width, bar.height = width, height
    for ax in (bar.x_axis, bar.y_axis):
        ax.delete = False
    return bar


def build_excel(path, days, kind, unit, col, pair, paircol):
    wb = openpyxl.Workbook()
    w0 = wb.active
    w0.title = "条件"
    w1 = wb.create_sheet("日ごとの稼働時間")
    w2 = wb.create_sheet("研究に使える時間")
    w3 = wb.create_sheet("期間別のまとめ")
    w4 = wb.create_sheet("グラフ")

    lines = [
        ("日中（9:00〜18:00）の室外機の稼働時間", FB),
        ("運転の判定（tools/operation_judge.py と同じ。push済みの「稼働状況の判定」と1時間ごとの判定が一致することを確認済み）", FN),
        ("① 1分ごと：排気温度 − 吸気温度（3方位の平均）の10分移動中央値が1.0℃以上なら「運転中」", F),
        ("② 1時間ごと：判定できた分が48分以上あり、そのうち80％以上が運転中なら「運転」（断続・停止・判定不能は数えない）", F),
        ("③ 日ごと：9〜17時台の9時間のうち「運転」の時間を数える。判定できた時間が7時間未満の日は空欄", F),
        ("④ 列の値：3台の平均（その日に値がある室外機だけで平均）", F),
        ("⑤ 研究に使える時間：平日に、組（芋緑化ありの室外機と同じ区画の列3）の2台がともに「運転」の時間。土日祝は空欄", F),
        ("空欄の主な理由：8-2の排気温度の欠測（8/14〜8/28）、3-2の排気温度の欠測（8/30〜9/7、9/14〜9/22ごろ）、"
         "列11のデータなし（9/15〜10/1）、7/29（設置当日）、10/13（14時ごろで記録終了）", FN),
        ("値はPythonで1分値から計算した。期間別のまとめは数式で集計している", FN),
    ]
    for i, (t, f) in enumerate(lines):
        w0.cell(row=1 + i, column=1, value=t).font = f
    w0.column_dimensions["A"].width = 140

    u_r0, u_r1 = table_sheet(w1, "日ごとの稼働時間［時間］（9〜17時台の9時間のうち「運転」の時間）", days, kind, [unit, col],
                             "灰色の行は土日祝。列3・列8・列11は3台の平均")
    p_r0, p_r1 = table_sheet(w2, "研究に使える時間［時間］（平日に2台とも「運転」の時間）", days, kind, [pair, paircol],
                             "土日祝は分析に使わないため空欄。列8・列11は3組の平均")

    # 期間別のまとめ（数式）
    w3["A1"] = "平日の平均［時間／9時間］と、研究に使える時間が7時間以上の平日の日数"
    w3["A1"].font = FB
    hdr = ["項目"] + [p[0] for p in PERIODS] + [f"{p[0]}：7時間以上の日数" for p in PERIODS]
    for j, h in enumerate(hdr):
        x = put(w3, 2, 1 + j, h, FH, fill=HEAD)
        w3.column_dimensions[x.column_letter].width = 16 if j == 0 else 11
    w3.row_dimensions[2].height = 30
    for j, (n, a, b) in enumerate(PERIODS):
        put(w3, 3, 2 + j, f"{a[5:].replace('-', '/')}〜{b[5:].replace('-', '/')}", FN)
    put(w3, 3, 1, "期間", FN)
    rows = [("稼働時間：" + c, "日ごとの稼働時間", 5 + len(unit.columns) + i) for i, c in enumerate(col.columns)]
    rows += [("研究に使える：" + c, "研究に使える時間", 5 + i) for i, c in enumerate(pair.columns)]
    rows += [("研究に使える：" + c + "（3組平均）", "研究に使える時間", 5 + len(pair.columns) + i) for i, c in enumerate(paircol.columns)]
    r = 4
    for label, sh, cidx in rows:
        put(w3, r, 1, label, F, align=LEFT)
        r0, r1 = (u_r0, u_r1) if sh == "日ごとの稼働時間" else (p_r0, p_r1)
        L = openpyxl.utils.get_column_letter(cidx)
        rng = f"'{sh}'!${L}${r0}:${L}${r1}"
        dts = f"'{sh}'!$A${r0}:$A${r1}"
        kd = f"'{sh}'!$D${r0}:$D${r1}"
        for j, (n, a, b) in enumerate(PERIODS):
            da = "DATE({},{},{})".format(*map(int, a.split("-")))
            db = "DATE({},{},{})".format(*map(int, b.split("-")))
            cond = f'{dts},">="&{da},{dts},"<="&{db},{kd},"平日"'
            put(w3, r, 2 + j, f'=IFERROR(AVERAGEIFS({rng},{cond}),"データなし")', fmt="0.0")
            if sh == "研究に使える時間" and "平均" not in label:
                put(w3, r, 6 + j, f'=IF(COUNTIFS({rng},">=0",{cond})=0,"データなし",COUNTIFS({rng},">=7",{cond}))')
            else:
                put(w3, r, 6 + j, "", F)
        r += 1

    # グラフ
    w4["A1"] = "グラフ（元データは「日ごとの稼働時間」「研究に使える時間」。灰色の帯は土日祝）"
    w4["A1"].font = FB
    ncol = len(unit.columns)
    w4.add_chart(combo(w1, u_r0, u_r1, [(5 + ncol + i, c) for i, c in enumerate(col.columns)],
                       [COLOR[c] for c in col.columns], "日中の稼働時間（3台の平均）",
                       band_col=5 + ncol + len(col.columns)), "A3")
    for k, (cname, us) in enumerate(COLS.items()):
        idx = [list(unit.columns).index(u) for u in us]
        w4.add_chart(combo(w1, u_r0, u_r1, [(5 + i, u) for i, u in zip(idx, us)], [UNIT_COLOR[u] for u in us],
                           f"{cname}：室外機ごとの日中の稼働時間", band_col=5 + ncol + len(col.columns),
                           height=6), f"A{20 + 13 * k}")
    npair = len(pair.columns)
    w4.add_chart(combo(w2, p_r0, p_r1, [(5 + npair + i, c) for i, c in enumerate(paircol.columns)],
                       [COLOR[c] for c in paircol.columns], "研究に使える時間（平日、3組の平均）",
                       band_col=5 + npair + len(paircol.columns)), "A60")
    wb.save(path)


def to_json(path, days, kind, unit, col, pair, paircol):
    def ser(fr):
        return {n: [val(v) for v in fr[n].values] for n in fr.columns}
    data = {
        "labels": [f"{d.month}/{d.day}" for d in days],
        "kind": [kind[d] for d in days],
        "band": [9 if kind[d] != "平日" else None for d in days],
        "unit": ser(unit), "col": ser(col), "pair": ser(pair), "paircol": ser(paircol),
    }
    wk = kind == "平日"
    summ = {}
    for name, fr in (("col", col), ("pair", pair), ("paircol", paircol)):
        summ[name] = {}
        for n in fr.columns:
            row = []
            for p, a, b in PERIODS:
                m = wk & (days >= a) & (days <= b)
                x = fr.loc[m, n]
                row.append({"mean": val(x.mean()) if x.notna().any() else None,
                            "days7": int((x >= 7).sum()) if x.notna().any() else None})
            summ[name][n] = row
    data["summary"] = summ
    data["periods"] = [p[0] for p in PERIODS]
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)


def main(src, xlsx, js):
    res = compute(src)
    build_excel(xlsx, *res)
    to_json(js, *res)


if __name__ == "__main__":
    main(*sys.argv[1:4])
